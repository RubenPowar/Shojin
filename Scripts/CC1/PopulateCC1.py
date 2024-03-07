import pandas as pd
from datetime import datetime
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm, Inches, Pt
from openpyxl import load_workbook
import docx, csv, logging, os
from dotenv import load_dotenv

import InsertTable as It
import GetMaps as Gm
import GetMapInfo as Gmi
import GetNearest as Gn


def start_logging():
    # Configure logging
    logging.basicConfig(filename='/Users/ruben/PycharmProjects/Shojin/files/Logs/CC1.log', format='%(message)s', level=logging.INFO)
    logging.info(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))


def log_file(filepath):
    logging.info(f"{filepath} was saved")


def format_currency(value):
    return "£{:,.2f}".format(value)


def format_percentage(value):
    return "{:.2f}%".format(value * 100)


# Find folders not labelled as compete (complete == CC1 completed)
def find_unfinished_sub_folders(folder):
    uf_sub_folders = []
    for folder_name in os.listdir(folder):
        full_path = os.path.join(folder, folder_name)
        # ---------- To be changed to colour of folder in Sharepoint
        if os.path.isdir(full_path) and 'unfinish' in folder_name.lower():
            uf_sub_folders.append(full_path)
    return uf_sub_folders


# Rename directory to reflect CC1 completed
def rename_directory(old_name):
    os.rename(old_name, old_name.replace("unfinish","finish"))


# Find completed appraisal excel
def find_appraisal_complete_excel(folder):
    ace_files = []
    for filename in os.listdir(folder):
        full_path = os.path.join(folder, filename)
        if os.path.isfile(full_path) and filename.lower().endswith('.xlsx') and 'appraisal-complete' in filename.lower():
            ace_files.append(full_path)
    return ace_files


# 0pen template and create new populated doc (temp filepaths)
def output_to_file(context, project_folder, category, template):
    # if category == 'Type1':
    #     template = DocxTemplate("/Users/ruben/PycharmProjects/Shojin/files/Logs/Templates/CC1-Template.docx")
    # else:
    #     template = DocxTemplate("/Users/ruben/PycharmProjects/Shojin/files/Logs/Templates/CC1-Template.docx")
    # print("Output _ :", context)
    project_title = context['project_title']
    template.render(context)
    project_title_us = project_title.replace(" ", "_")
    new_file = f"{project_folder}/Generated/CC1-{project_title_us}.docx"
    # print("Saving File")
    template.save(new_file)
    log_file(new_file)


def get_context_data_from_coversheet(coversheet):
    context = {
        'project_title': coversheet['D2'].value,
        'lead': coversheet['D3'].value,
        'introducer': coversheet['D4'].value,
        'developer': coversheet['D5'].value,
        'drawdown_date': coversheet['D6'].value.strftime('%B %Y'),
        'anticipated_term': coversheet['D7'].value,
        'type': coversheet['D8'].value,
        'combined_loan': format_currency(coversheet['D9'].value),
        'ltv': format_percentage(coversheet['D10'].value),
        'net_facility': format_currency(coversheet['D11'].value),
        'borrower_equity': format_currency(coversheet['D12'].value),
        'combined_net_release': format_percentage(coversheet['D13'].value),
        'interest_rate': format_percentage(coversheet['D15'].value),
        'address': coversheet['D17'].value,
        'market_value': format_currency(coversheet['D23'].value),
        'am_fee': coversheet['D32'].value,
        'professional_costs': format_currency(coversheet['D33'].value),
        'interest_reserve': coversheet['D34'].value,

        'today_date': datetime.today().strftime("%d %b, %Y"),
        # Other formatted fields omitted for brevity
    }
    return context


def get_tables_data(tablesheet, context):
    # Assuming get_table_data is defined elsewhere and works as intended
    for table in tablesheet.tables.values():
        context = It.get_table_data(table.ref, tablesheet[table.ref], table.name, context)
    return context


def generate_context(wb, template, project_folder, gmaps_key):
    coversheet = wb["CC1 Coversheet"]
    context = get_context_data_from_coversheet(coversheet)
    #For Template CC1
    # Get variables from Coversheet
    # coversheet = wb["Coversheet"]
    # project_title = coversheet['B2'].value #ws.loc[1, 1]
    # my_name = coversheet['B3'].value
    # introducer = coversheet['B4'].value
    # developer = coversheet['B5'].value
    # combined_loan = coversheet['B6'].value
    # drawdown_date = coversheet['B7'].value
    # anticipated_term = coversheet['B8'].value
    # postcode = coversheet['B9'].value
    # today_date = datetime.today().strftime("%d %b, %Y")
    # context = {'project_title': project_title + ', ' + postcode, 'my_name': my_name, 'introducer': introducer,
    #               'developer': developer, 'combined_loan': combined_loan, 'drawdown_date': drawdown_date,
    #               'anticipated_term': anticipated_term, 'today_date': today_date}

    # For template 2 (Sam's)
    # project_title = coversheet['D2'].value  # ws.loc[1, 1]
    # lead = coversheet['D3'].value
    # introducer = coversheet['D4'].value
    # developer = coversheet['D5'].value
    # drawdown_date = coversheet['D6'].value.strftime('%B %Y')
    # anticipated_term = coversheet['D7'].value
    # type = coversheet['D8'].value
    # combined_loan = format_currency(coversheet['D9'].value)
    # ltv = format_percentage(coversheet['D10'].value * 100)
    # net_facility = format_currency(coversheet['D11'].value)
    # borrower_equity = format_currency(coversheet['D12'].value)
    # combined_net_release = format_percentage(coversheet['D13'].value * 100)
    # interest_rate = format_percentage(coversheet['D15'].value * 100)
    #
    # address = coversheet['D17'].value
    # market_value = coversheet['D23'].value
    #
    # # Shojin Terms
    # am_fee = coversheet['D32'].value
    # professional_costs = coversheet['D33'].value
    # interest_reserve = coversheet['D34'].value
    #
    #
    # today_date = datetime.today().strftime("%d %b, %Y")
    # context = {'project_title': project_title, 'lead': lead, 'introducer': introducer,
    #            'developer': developer, 'drawdown_date': drawdown_date, 'anticipated_term': anticipated_term,
    #            'type': type, 'combined_loan': combined_loan, 'ltv': ltv, 'net_facility': net_facility,
    #            'borrower_equity': borrower_equity, 'combined_net_release': combined_net_release,
    #            'interest_rate': interest_rate, 'address': address, 'market_value': market_value, 'am_fee': am_fee,
    #            'professional_costs': professional_costs, 'interest_reserve': interest_reserve}

    # Get tables from Tables sheet
    tablesheet = wb["CC1 Tables"]
    context = get_tables_data(tablesheet, context)
    # index = 1
    # for table in tablesheet.tables.values():
    #     name = table.name
    #     ref = table.ref
    #     cell_range = tablesheet[ref]
    #     # References InsertTable.py script
    #     context = It.get_table_data(ref, cell_range, name, context)
    #     index += 1
    # print("project folder = ", project_folder)
    filepath = os.path.join(project_folder, "Generated/")

    # Generate maps images and save to file
    filenames = Gm.generate_maps_images(context['address'], context['project_title'], filepath, gmaps_key)
    coordinates = Gm.get_coordinates(context['address'])

    # Generate text for amenities and transport links
    amenities_text = Gmi.generate_maps_info(coordinates, filepath, gmaps_key)
    transport_text = Gn.generate_nearest_info(coordinates)

    # Insert selected maps images from list or filepaths: filenames[]
    context.update({
        'map_1': InlineImage(template, image_descriptor=filenames[1], width=Mm(170), height=Mm(170)),
        'map_2': InlineImage(template, image_descriptor=filenames[2], width=Mm(170), height=Mm(170)),
        'transport_text': transport_text,
        'amenities_text': amenities_text
    })
    return context


# Main finds values in appraisal
def generate_file(project_folder, appraisal, gmaps_key, template):
    # template = DocxTemplate("/Users/ruben/PycharmProjects/Shojin/files/CC1/Sources/CC1-Bridge-Template.docx")
    wb = load_workbook(appraisal, data_only=True)
    context = generate_context(wb, template, project_folder, gmaps_key)
    category = 'Type1' # type of proposal (determines style of CC1)
    output_to_file(context, project_folder, category, template)


def get_excel_file(folder):
    appraisal_complete_excel_files = find_appraisal_complete_excel(folder)
    # print(folder)
    # print(appraisal_complete_excel_files)
    # main(folder, appraisal_complete_excel_files[0], )


if __name__ == '__main__':
    load_dotenv()
    sharepoint_password = os.getenv('SHAREPOINT-PASSWORD')
    sharepoint_user = os.getenv('SHAREPOINT-USER')
    gmaps_api_key = os.getenv('GMAPS_API_KEY')
    # print(sharepoint_user, sharepoint_password)
    start_logging()
    # root_folder = "/Users/ruben/PycharmProjects/Shojin/files/Sharepoint/"

    # Declare files to read
    # unfinished_sub_folders = find_unfinished_sub_folders(root_folder)
    CC1_project_folder = "/Users/ruben/PycharmProjects/Shojin/files/CC1/"
    CC1_template = DocxTemplate("/Users/ruben/PycharmProjects/Shojin/files/CC1/Sources/CC1-Bridge-Template.docx")
    sources_folder = CC1_project_folder + 'Sources/'
    xl_source = 'Bridge-XLS.xlsm'
    generate_file(CC1_project_folder, sources_folder + xl_source, gmaps_api_key, CC1_template)


    # for sub_folder in unfinished_sub_folders:
    #     get_excel_file(sub_folder)
    #     appraisal_complete_excel_files = find_appraisal_complete_excel(sub_folder)
    #     main(sub_folder, appraisal_complete_excel_files[0])
    #     rename_directory(sub_folder)


    # sys.exit(main(sys.argv[1:]))


# Sources:
#     https://www.youtube.com/watch?v=H8Ars15wGR
#     https://github.com/thepycoach/automation/blob/main/Word/en-word.py


